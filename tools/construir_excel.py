"""Builds the Part A workbook exactly as specified in the workshop statement.

Usage:
    python tools/construir_excel.py SALIDA.xlsx [valores.json]

Without a JSON file the workbook only contains formulas (used as input for the
LibreOffice computation in tools/calcular_libreoffice.py). With the JSON file
the 100 replications of the data table and the "Politicas" sheet are filled in
with the values produced by the spreadsheet itself, and every formula cell
receives its cached value (tools/inyectar_valores.py).

Formulas are stored in Excel's internal (English) syntax; Spanish Excel shows
them translated automatically (RANDBETWEEN -> ALEATORIO.ENTRE, ";" separator).

Visual direction ("achromatic ledger"): no coloured header fills; ink text on a
blue-grey ledger tone, hairline row rules, money in bold only where it is a
result, zeros shown as a dash, and one chart hue (#1D5FBF) reserved for the main
data element of each chart.
"""

import json
import sys

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference, ScatterChart, Series
from openpyxl.chart.axis import ChartLines
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.marker import DataPoint, Marker
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText, Text
from openpyxl.chart.title import Title
from openpyxl.chart.data_source import NumFmt
from openpyxl.chart.trendline import Trendline, TrendlineLabel
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import (CharacterProperties, Paragraph, ParagraphProperties,
                                   RegularTextRun)
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.formula import DataTableFormula

FIRST, LAST = 6, 318            # day 1 .. day 313
REP_FIRST, REP_LAST = 7, 106    # 100 replications

# --- Design tokens -----------------------------------------------------------
FONT = "Aptos"
INK, MUTED, RULE, CONTROL = "141414", "6B6B6B", "E6E6E6", "BFBFBF"
LEDGER = "F2F4F7"          # blue-grey ledger ground for headers and result blocks
HUE, HUE_LIGHT = "1D5FBF", "9DB7E3"
GREY_SERIES = "8A8A8A"
NEG = "C4242A"

F_BODY = Font(name=FONT, size=10, color=INK)
F_MUTED = Font(name=FONT, size=9, color=MUTED)
F_HEAD = Font(name=FONT, size=9, color="3D3D3D", bold=True)
F_TITLE = Font(name=FONT, size=16, color=INK, bold=True)
F_SECTION = Font(name=FONT, size=11, color=INK, bold=True)
F_RESULT = Font(name=FONT, size=11, color=INK, bold=True)
F_INPUT = Font(name=FONT, size=10, color=HUE, bold=True)

FILL_LEDGER = PatternFill("solid", fgColor=LEDGER)
FILL_INPUT = PatternFill("solid", fgColor="EAF0FA")
FILL_WHITE = PatternFill("solid", fgColor="FFFFFF")
HAIR = Side(style="thin", color=RULE)
UNDER_HEAD = Side(style="thin", color=CONTROL)
ROW_RULE = Border(bottom=HAIR)
HEAD_RULE = Border(bottom=UNDER_HEAD)

WRAP_LEFT = Alignment(wrap_text=True, vertical="bottom", horizontal="left")
WRAP_RIGHT = Alignment(wrap_text=True, vertical="bottom", horizontal="right")

MONEY = '"$"#,##0;[Red]-"$"#,##0;"–"'
MILLIONS = '"$"#,##0.0,,"M"'
RATIO = '0.0000;-0.0000;"–"'
DATE = "[$-240A]ddd dd/mm/yyyy"     # es-CO locale id: Spanish day names in any Excel


# --- Cell helpers -------------------------------------------------------------
def put(ws, ref, value, font=F_BODY, fmt=None, fill=None, border=None, align=None):
    c = ws[ref] if isinstance(ref, str) else ws.cell(*ref)
    c.value = value
    c.font = font
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if border:
        c.border = border
    if align:
        c.alignment = align
    return c


def header(ws, row, col, text, width=None, right=False):
    c = put(ws, (row, col), text, F_HEAD, fill=FILL_LEDGER, border=HEAD_RULE,
            align=WRAP_RIGHT if right else WRAP_LEFT)
    if width:
        ws.column_dimensions[c.column_letter].width = width
    return c


def title_block(ws, title, subtitle):
    put(ws, "A1", title, F_TITLE)
    put(ws, "A2", subtitle, F_MUTED)
    ws.row_dimensions[1].height = 24
    ws.sheet_view.showGridLines = False


# --- Chart helpers --------------------------------------------------------------
def _cp(size, color=INK, bold=False):
    return CharacterProperties(sz=size, b=bold, solidFill=color, latin=None)


def rich_title(text, size=1100, color=INK, bold=True):
    cp = _cp(size, color, bold)
    para = Paragraph(pPr=ParagraphProperties(defRPr=cp), r=[RegularTextRun(rPr=cp, t=text)])
    return Title(tx=Text(rich=RichText(p=[para])), overlay=False)


def tick_text(size=800, color=MUTED):
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=_cp(size, color)),
                                 endParaRPr=_cp(size, color))])


def style_axes(chart, x_title, y_title):
    chart.x_axis.title = rich_title(x_title, 900, MUTED, False)
    chart.y_axis.title = rich_title(y_title, 900, MUTED, False)
    for ax in (chart.x_axis, chart.y_axis):
        ax.delete = False
        ax.txPr = tick_text()
        ax.spPr = GraphicalProperties(ln=LineProperties(solidFill=CONTROL))
    chart.y_axis.majorGridlines = ChartLines(
        spPr=GraphicalProperties(ln=LineProperties(solidFill="EDEDED")))
    chart.x_axis.majorGridlines = None
    chart.plot_area.graphicalProperties = GraphicalProperties(noFill=True)
    chart.graphical_properties = GraphicalProperties(
        solidFill="FFFFFF", ln=LineProperties(solidFill=RULE))


def legend_bottom(chart):
    chart.legend.position = "t"   # top: never collides with the x-axis title
    chart.legend.txPr = tick_text(800, "3D3D3D")


# --- Sheet: Parametros -----------------------------------------------------------
def build_parametros(wb, policy):
    ws = wb.active
    ws.title = "Parametros"
    title_block(ws, "Parámetros del modelo",
                "Taller 2 · Simulación Montecarlo de un sistema de inventarios · "
                "Juan Andrés Young Hoyos (1000337497). Texto azul: política (s, Q) elegida.")
    header(ws, 3, 1, "Parámetro", 40)
    header(ws, 3, 2, "Valor", 14, right=True)
    header(ws, 3, 3, "Nombre definido", 17)
    ws.row_dimensions[3].height = 30
    ws.column_dimensions["D"].width = 3
    params = [
        (4, "Punto de reorden s", policy[0], "PuntoReorden", "0"),
        (5, "Cantidad a pedir Q", policy[1], "CantPedido", "0"),
        (6, "Precio de venta (por unidad)", 680000, "PrecioVenta", MONEY),
        (7, "Costo de compra (por unidad recibida)", 600000, "CostoUnitario", MONEY),
        (8, "Costo por pedido recibido", 40000, "CostoPedido", MONEY),
        (9, "Almacenamiento por unidad-día (inv. promedio)", 80, "CostoAlm", MONEY),
        (10, "Costo fijo mensual (primer día hábil del mes)", 1000000, "CostoFijo", MONEY),
        (11, "Año simulado", 2026, "AnioSim", "0"),
    ]
    for row, label, value, name, fmt in params:
        is_policy = row in (4, 5)
        put(ws, (row, 1), label, border=ROW_RULE)
        put(ws, (row, 2), value, F_INPUT if is_policy else F_BODY, fmt,
            FILL_INPUT if is_policy else None, ROW_RULE)
        put(ws, (row, 3), name, F_MUTED, border=ROW_RULE)
        wb.defined_names[name] = DefinedName(name, attr_text=f"Parametros!$B${row}")

    # Step A2: empirical distribution (inverse transform with VLOOKUP)
    for col, (text, width) in enumerate([("Lím. inf. acumulado", 12), ("Demanda", 10),
                                         ("Probabilidad", 12), ("Acumulada", 12)], start=5):
        header(ws, 3, col, text, width, right=True)
    probs = [0.20, 0.05, 0.40, 0.25, 0.10]
    for i, (d, p) in enumerate(zip(range(3, 8), probs)):
        r = 4 + i
        put(ws, (r, 5), 0 if i == 0 else f"=H{r - 1}", fmt="0.00", border=ROW_RULE)
        put(ws, (r, 6), d, border=ROW_RULE)
        put(ws, (r, 7), p, fmt="0.00", border=ROW_RULE)
        put(ws, (r, 8), f"=G{r}" if i == 0 else f"=H{r - 1}+G{r}", fmt="0.00",
            border=ROW_RULE)
    put(ws, "F9", "Suma", F_HEAD)
    put(ws, "G9", "=SUM(G4:G8)", F_RESULT, "0.00")
    put(ws, "H9", '=IF(ROUND(G9,10)=1,"Correcta","Revisar")', F_MUTED)
    wb.defined_names["TablaEmpirica"] = DefinedName("TablaEmpirica",
                                                    attr_text="Parametros!$E$4:$F$8")
    put(ws, "E10", "TablaEmpirica = E4:F8 · uniforme: ALEATORIO.ENTRE(3;7) · empírica: "
                   "BUSCARV(ALEATORIO();TablaEmpirica;2;VERDADERO)", F_MUTED)

    # Part A question 2: theoretical mean and variance of daily demand
    put(ws, "A13", "Demanda diaria: distribución, media y varianza teóricas", F_SECTION)
    for col, text in enumerate(["Demanda d", "P uniforme", "P empírica"], start=1):
        header(ws, 14, col, text, right=col > 1)
    for i, d in enumerate(range(3, 8)):
        r = 15 + i
        put(ws, (r, 1), d, border=ROW_RULE)
        put(ws, (r, 2), 0.2, fmt="0.00", border=ROW_RULE)
        put(ws, (r, 3), f"=G{4 + i}", fmt="0.00", border=ROW_RULE)
    for col, text in enumerate(["Distribución", "Media E[D]", "Varianza Var[D]",
                                "Desv. estándar"], start=1):
        header(ws, 21, col, text, right=col > 1)
    for r, name, pcol in ((22, "Uniforme discreta", "B"), (23, "Empírica", "C")):
        put(ws, (r, 1), name, border=ROW_RULE)
        put(ws, (r, 2), f"=SUMPRODUCT($A$15:$A$19,{pcol}15:{pcol}19)", F_RESULT, "0.00",
            border=ROW_RULE)
        put(ws, (r, 3), f"=SUMPRODUCT(($A$15:$A$19-B{r})^2,{pcol}15:{pcol}19)", F_RESULT,
            "0.00", border=ROW_RULE)
        put(ws, (r, 4), f"=SQRT(C{r})", fmt="0.000", border=ROW_RULE)
    ws.column_dimensions["D"].width = 14

    chart = BarChart()
    chart.type = "col"
    chart.grouping = "clustered"
    chart.gapWidth = 60
    chart.overlap = -10
    chart.title = rich_title("Misma media (5), distinta dispersión")
    for col, color in ((2, GREY_SERIES), (3, HUE)):
        chart.add_data(Reference(ws, min_col=col, min_row=14, max_row=19),
                       titles_from_data=True)
        chart.series[-1].graphicalProperties = GraphicalProperties(solidFill=color,
                                                                   ln=LineProperties(noFill=True))
    chart.set_categories(Reference(ws, min_col=1, min_row=15, max_row=19))
    style_axes(chart, "Demanda diaria (unidades)", "Probabilidad")
    chart.y_axis.numFmt = "0.00"
    chart.y_axis.scaling.min = 0
    legend_bottom(chart)
    chart.height, chart.width = 7.5, 15
    ws.add_chart(chart, "F13")

    put(ws, "A26", "Cálculo", F_SECTION)
    put(ws, "A27", "Archivo > Opciones > Fórmulas > «Automático excepto en las tablas de "
                   "datos». F9 genera 100 réplicas nuevas.", F_MUTED)
    put(ws, "A28", "Los estadísticos del informe corresponden a las 100 réplicas guardadas "
                   "en este archivo.", F_MUTED)


# --- Sheets: Uniforme / Empirica -----------------------------------------------
DAILY_HEADERS = [
    ("Día", 6), ("Fecha", 15), ("Mes", 5), ("Inicio de mes (1/0)", 8),
    ("Unidades recibidas", 10), ("Inventario inicial", 10), ("Demanda", 9),
    ("Ventas", 8), ("Ventas perdidas", 9), ("Inventario final", 10),
    ("Día con faltante (1/0)", 9), ("Pedido al cierre (1/0)", 9),
    ("Ingresos", 13), ("Costo de compra", 13), ("Costo de pedido", 11),
    ("Costo de almacena-miento", 12), ("Costo fijo", 12), ("Utilidad del día", 14),
]
MONEY_COLS = range(13, 19)


def daily_row(r, demand_formula):
    """Formulas of row r (6 = day 1, 7..318 = general row) exactly as in the statement."""
    if r == FIRST:
        return ["1", "=DATE(AnioSim,1,1)", "=MONTH(B6)", "1", "=CantPedido", "=E6",
                demand_formula, "=MIN(F6,G6)", "=G6-H6", "=F6-H6", "=IF(I6>0,1,0)",
                "=IF(J6<PuntoReorden,1,0)", "=PrecioVenta*H6", "=CostoUnitario*E6",
                "=IF(E6>0,CostoPedido,0)", "=CostoAlm*(F6+J6)/2", "=CostoFijo*D6",
                "=M6-N6-O6-P6-Q6"]
    p = r - 1
    return [f"=A{p}+1", f"=B{p}+IF(WEEKDAY(B{p},2)=6,2,1)", f"=MONTH(B{r})",
            f"=IF(C{r}<>C{p},1,0)", f"=IF(L{p}=1,CantPedido,0)", f"=J{p}+E{r}",
            demand_formula, f"=MIN(F{r},G{r})", f"=G{r}-H{r}", f"=F{r}-H{r}",
            f"=IF(I{r}>0,1,0)", f"=IF(J{r}<PuntoReorden,1,0)", f"=PrecioVenta*H{r}",
            f"=CostoUnitario*E{r}", f"=IF(E{r}>0,CostoPedido,0)",
            f"=CostoAlm*(F{r}+J{r})/2", f"=CostoFijo*D{r}",
            f"=M{r}-N{r}-O{r}-P{r}-Q{r}"]


def histogram_chart(ws, label):
    chart = BarChart()
    chart.type = "col"
    chart.gapWidth = 0      # adjacent bars (ancho de intervalo 0 %)
    chart.title = rich_title(f"Histograma de la ganancia anual · demanda {label}")
    chart.add_data(Reference(ws, min_col=32, min_row=17, max_row=26), titles_from_data=False)
    chart.set_categories(Reference(ws, min_col=31, min_row=17, max_row=26))
    s = chart.series[0]
    s.graphicalProperties = GraphicalProperties(solidFill=HUE,
                                                ln=LineProperties(solidFill="FFFFFF", w=12700))
    s.dLbls = DataLabelList()
    s.dLbls.showVal = True
    for flag in ("showSerName", "showCatName", "showLegendKey", "showPercent"):
        setattr(s.dLbls, flag, False)
    s.dLbls.txPr = tick_text(800, INK)
    style_axes(chart, "Marca de clase de la ganancia anual (millones de COP)",
               "Frecuencia (réplicas)")
    chart.x_axis.numFmt = "0.0"
    chart.y_axis.scaling.min = 0
    chart.legend = None
    chart.height, chart.width = 8.5, 17
    return chart


def inventory_chart(ws, label):
    chart = LineChart()
    last = FIRST + 59   # first 60 working days: the (s, Q) cycle is readable
    chart.title = rich_title(f"Inventario final, primeros 60 días · réplica actual ({label})")
    chart.add_data(Reference(ws, min_col=10, min_row=5, max_row=last), titles_from_data=True)
    chart.add_data(Reference(ws, min_col=19, min_row=5, max_row=last), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=FIRST, max_row=last))
    inv, s_line = chart.series
    inv.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill=HUE, w=15875))
    s_line.graphicalProperties = GraphicalProperties(
        ln=LineProperties(solidFill=INK, w=12700, prstDash="dash"))
    for ser in chart.series:
        ser.marker = Marker(symbol="none")
        ser.smooth = False
    style_axes(chart, "Día hábil de 2026", "Unidades")
    chart.x_axis.tickLblSkip = 5
    chart.x_axis.tickMarkSkip = 5
    chart.y_axis.scaling.min = 0
    legend_bottom(chart)
    chart.height, chart.width = 8.5, 17
    return chart


def leftover_chart(ws, label):
    chart = ScatterChart()
    chart.title = rich_title(f"Ganancia anual vs. inventario sobrante al 31/12 ({label})")
    xs = Reference(ws, min_col=27, min_row=REP_FIRST, max_row=REP_LAST)
    ys = Reference(ws, min_col=24, min_row=REP_FIRST, max_row=REP_LAST)
    ser = Series(ys, xs, title="Réplicas")
    ser.marker = Marker(symbol="circle", size=5)
    ser.marker.graphicalProperties = GraphicalProperties(solidFill=HUE,
                                                         ln=LineProperties(noFill=True))
    ser.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))
    ser.trendline = Trendline(
        trendlineType="linear", dispRSqr=True, dispEq=False,
        trendlineLbl=TrendlineLabel(numFmt=NumFmt(formatCode="0.000", sourceLinked=False),
                                    txPr=tick_text(900, INK)),
        spPr=GraphicalProperties(ln=LineProperties(solidFill=INK, w=12700)))
    chart.series.append(ser)
    style_axes(chart, "Inventario final al 31/12 (unidades, sin valor residual)",
               "Ganancia anual (COP)")
    chart.y_axis.numFmt = MILLIONS
    chart.x_axis.scaling.min = 0
    chart.legend = None
    chart.height, chart.width = 8.5, 17
    return chart


def build_simulation(wb, title, demand_formula, label, replications):
    ws = wb.create_sheet(title)
    title_block(ws, f"Un año simulado · demanda {label}",
                "313 días hábiles de 2026 (lunes a sábado). Fila 6 = día 1; la fila 7 se "
                "copia hasta la 318. Cada recálculo (F9) es una réplica distinta.")
    put(ws, "A3", "Demanda (columna G): " + (
        "ALEATORIO.ENTRE(3;7)" if title == "Uniforme"
        else "BUSCARV(ALEATORIO();TablaEmpirica;2;VERDADERO), transformada inversa"),
        F_MUTED)
    for col, (text, width) in enumerate(DAILY_HEADERS, start=1):
        header(ws, 5, col, text, width, right=col not in (2,))
    ws.row_dimensions[5].height = 40

    for r in range(FIRST, LAST + 1):
        for col, f in enumerate(daily_row(r, demand_formula), start=1):
            fmt = DATE if col == 2 else (MONEY if col in MONEY_COLS else None)
            font = F_BODY
            put(ws, (r, col), int(f) if f.isdigit() else f, font, fmt, border=ROW_RULE)
        put(ws, (r, 19), "=PuntoReorden", F_MUTED, border=ROW_RULE,
            align=Alignment(horizontal="center"))
    header(ws, 5, 19, "Punto de reorden s", 9, right=True)

    # Step A4: replication summary (U6:U8, plus the 31/12 inventory in U9)
    put(ws, "T4", "Resumen de la réplica", F_SECTION, align=Alignment(indent=1))
    ws.column_dimensions["T"].width = 33
    ws.column_dimensions["U"].width = 15
    summary = [(6, "Ganancia anual", "=SUM(R6:R318)", MONEY),
               (7, "Fracción de días con faltante", "=SUM(K6:K318)/COUNT(A6:A318)", RATIO),
               (8, "Ventas perdidas promedio por día", "=AVERAGE(I6:I318)", RATIO),
               (9, "Inventario final al 31/12/2026", "=J318", "0")]
    for r, text, f, fmt in summary:
        put(ws, (r, 20), text, border=ROW_RULE, align=Alignment(indent=1))
        put(ws, (r, 21), f, F_RESULT, fmt, FILL_LEDGER, ROW_RULE)
    put(ws, "T10", "Comprobación (B318, A318, suma de D)", F_MUTED, align=Alignment(indent=1))
    put(ws, "U10", '=IF(AND(B318=DATE(AnioSim,12,31),A318=313,SUM(D6:D318)=12),'
                   '"Correcta","Revisar")', F_MUTED)

    # Step A5: data table with 100 replications (column input cell W2, empty)
    put(ws, "W2", None)
    put(ws, "W1", "W2 = celda de entrada de la tabla de datos (debe quedar vacía).", F_MUTED)
    for col, (text, width) in enumerate([("Réplica", 8), ("Ganancia anual", 15),
                                         ("Fracción días con faltante", 11),
                                         ("Ventas perdidas por día", 11),
                                         ("Inventario final 31/12", 10)], start=23):
        header(ws, 5, col, text, width, right=True)
    links = (("X6", "=U6", MONEY), ("Y6", "=U7", RATIO), ("Z6", "=U8", RATIO),
             ("AA6", "=U9", "0"))
    for ref, f, fmt in links:
        put(ws, ref, f, F_MUTED, fmt, border=ROW_RULE)
    for i in range(100):
        r = REP_FIRST + i
        put(ws, (r, 23), i + 1, F_MUTED, border=ROW_RULE)
        for col, fmt in ((24, MONEY), (25, RATIO), (26, RATIO), (27, "0")):
            put(ws, (r, col), None, fmt=fmt, border=ROW_RULE)
    # {=TABLE(,W2)}: Data > What-If Analysis > Data Table, column input cell W2
    ws["X7"] = DataTableFormula(ref="X7:AA106", r1="W2", dt2D=False, dtr=False)
    if replications:
        for i, row in enumerate(replications):
            for j, v in enumerate(row):
                if (i, j) != (0, 0):   # X7 holds the TABLE formula
                    ws.cell(REP_FIRST + i, 24 + j).value = v

    # Step A6: statistics
    put(ws, "AB4", "Estadísticos de las 100 réplicas", F_SECTION, align=Alignment(indent=1))
    ws.column_dimensions["AB"].width = 36
    for col in ("AC", "AD", "AE", "AF"):
        ws.column_dimensions[col].width = 15
    stats = [(5, "Ganancia promedio", "=AVERAGE(X7:X106)", MONEY),
             (6, "Desviación estándar de la ganancia", "=_xlfn.STDEV.S(X7:X106)", MONEY),
             (7, "Probabilidad de ventas perdidas", "=AVERAGE(Y7:Y106)", RATIO),
             (8, "Ventas perdidas promedio por día", "=AVERAGE(Z7:Z106)", RATIO),
             (9, "Error estándar de la media", "=AC6/SQRT(COUNT(X7:X106))", MONEY)]
    for r, text, f, fmt in stats:
        put(ws, (r, 28), text, border=ROW_RULE, align=Alignment(indent=1))
        put(ws, (r, 29), f, F_RESULT if r == 5 else F_BODY, fmt, FILL_LEDGER, ROW_RULE)

    # Step A7: histogram with 10 equal-width classes
    put(ws, "AB11", "Histograma: 10 clases de igual ancho", F_SECTION, align=Alignment(indent=1))
    for r, text, f in ((12, "Mínimo", "=MIN(X7:X106)"), (13, "Máximo", "=MAX(X7:X106)"),
                       (14, "Ancho de clase", "=(AC13-AC12)/10")):
        put(ws, (r, 28), text, border=ROW_RULE, align=Alignment(indent=1))
        put(ws, (r, 29), f, fmt=MONEY, border=ROW_RULE)
    for col, text in enumerate(["Clase k", "Límite inferior", "Límite superior",
                                "Marca (millones)", "Frecuencia"], start=28):
        header(ws, 16, col, text, right=col > 28)
        ws.cell(16, col).alignment = Alignment(horizontal="right" if col > 28 else "left")
    ws.row_dimensions[16].height = 15   # keep the daily table rows uniform
    for k in range(1, 11):
        r = 16 + k
        op = "<=" if k == 10 else "<"
        put(ws, (r, 28), k, F_MUTED, border=ROW_RULE, align=Alignment(horizontal="center"))
        put(ws, (r, 29), f"=$AC$12+(AB{r}-1)*$AC$14", fmt=MONEY, border=ROW_RULE)
        put(ws, (r, 30), f"=AC{r}+$AC$14", fmt=MONEY, border=ROW_RULE)
        put(ws, (r, 31), f"=(AC{r}+AD{r})/2/1000000", fmt="0.00", border=ROW_RULE)
        put(ws, (r, 32), f'=COUNTIFS($X$7:$X$106,">="&AC{r},$X$7:$X$106,"{op}"&AD{r})',
            F_RESULT, border=ROW_RULE)
    put(ws, "AE27", "Total", F_HEAD, align=WRAP_RIGHT)
    put(ws, "AF27", "=SUM(AF17:AF26)", F_RESULT)

    # Part A question 4: profit vs inventory left on 31/12
    put(ws, "AB29", "Ganancia vs. inventario sobrante al 31/12", F_SECTION,
        align=Alignment(indent=1))
    for r, text, f, fmt in (
            (30, "Inventario final promedio (unidades)", "=AVERAGE(AA7:AA106)", "0.00"),
            (31, "Correlación ganancia – inventario final", "=CORREL(X7:X106,AA7:AA106)",
             "0.000"),
            (32, "Pendiente ($ por unidad sobrante)", "=SLOPE(X7:X106,AA7:AA106)", MONEY),
            (33, "R² (variabilidad explicada)", "=AC31^2", "0.000")):
        put(ws, (r, 28), text, border=ROW_RULE, align=Alignment(indent=1))
        put(ws, (r, 29), f, fmt=fmt, border=ROW_RULE)

    ws.add_chart(histogram_chart(ws, label), "AH4")
    ws.add_chart(inventory_chart(ws, label), "AH22")
    ws.add_chart(leftover_chart(ws, label), "AH40")
    return ws


# --- Sheet: Politicas --------------------------------------------------------------
def build_politicas(wb, policies, recommended):
    ws = wb.create_sheet("Politicas")
    title_block(ws, "Comparación de políticas (s, Q)",
                "100 réplicas por combinación y caso de demanda (Paso A9): se cambian s y Q "
                "en Parametros, se recalcula y se guardan los estadísticos como valores.")
    heads = [("Demanda", 12), ("s", 6), ("Q", 6), ("Ganancia promedio", 16),
             ("Desv. estándar", 15), ("Prob. ventas perdidas", 12),
             ("Ventas perdidas por día", 12), ("Observación", 36)]
    for col, (text, width) in enumerate(heads, start=1):
        header(ws, 4, col, text, width, right=1 < col < 8)
    ws.row_dimensions[4].height = 44
    r = 5
    for case in ("Uniforme", "Empírica"):
        for row in policies.get(case, []):
            rec = (row["s"], row["Q"]) == tuple(recommended)
            vals = [(case, None), (row["s"], "0"), (row["Q"], "0"), (row["mean"], MONEY),
                    (row["std"], MONEY), (row["p_lost"], RATIO), (row["lost_day"], RATIO)]
            for col, (v, fmt) in enumerate(vals, start=1):
                font = F_RESULT if (rec and col == 4) else (F_INPUT if rec and col in (2, 3)
                                                            else F_BODY)
                put(ws, (r, col), v, font, fmt, FILL_INPUT if rec else None, ROW_RULE)
            put(ws, (r, 8), "Recomendada (cargada en Parametros)" if rec else None,
                F_INPUT if rec else F_MUTED, fill=FILL_INPUT if rec else None,
                border=ROW_RULE)
            r += 1
    last_row = r - 1

    put(ws, (r + 1, 1), "Criterio de selección", F_SECTION)
    notes = ["1. Servicio: cero ventas perdidas con ambas demandas (s ≥ 7 y Q ≥ 7 lo "
             "garantizan; ver informe, pregunta 5).",
             "2. Ganancia: entre esas políticas, la mayor ganancia promedio de los dos casos.",
             "3. Riesgo: si las diferencias son menores que el error de Montecarlo, la de "
             "menor desviación estándar."]
    for i, text in enumerate(notes):
        put(ws, (r + 2 + i, 1), text, F_MUTED)

    # Chart data block: one row per policy, both demand cases side by side
    unif = policies.get("Uniforme", [])
    emp = {(p["s"], p["Q"]): p for p in policies.get("Empírica", [])}
    header(ws, 4, 10, "Política", 14)
    header(ws, 4, 11, "Uniforme", 15, right=True)
    header(ws, 4, 12, "Empírica", 15, right=True)
    rec_idx = None
    for i, p in enumerate(unif):
        rr = 5 + i
        put(ws, (rr, 10), f"s={p['s']}, Q={p['Q']}", border=ROW_RULE)
        put(ws, (rr, 11), p["mean"], fmt=MONEY, border=ROW_RULE)
        put(ws, (rr, 12), emp[(p["s"], p["Q"])]["mean"], fmt=MONEY, border=ROW_RULE)
        if (p["s"], p["Q"]) == tuple(recommended):
            rec_idx = i
    if unif:
        n = len(unif)
        chart = BarChart()
        chart.type = "bar"
        chart.grouping = "clustered"
        chart.gapWidth = 50
        chart.title = rich_title("Ganancia promedio por política (recomendada en azul)")
        for col, color, light in ((11, INK, HUE), (12, GREY_SERIES, HUE_LIGHT)):
            chart.add_data(Reference(ws, min_col=col, min_row=4, max_row=4 + n),
                           titles_from_data=True)
            ser = chart.series[-1]
            ser.graphicalProperties = GraphicalProperties(solidFill=color,
                                                          ln=LineProperties(noFill=True))
            if rec_idx is not None:
                ser.dPt = [DataPoint(idx=rec_idx, invertIfNegative=False,
                                     spPr=GraphicalProperties(solidFill=light,
                                                              ln=LineProperties(noFill=True)))]
        chart.set_categories(Reference(ws, min_col=10, min_row=5, max_row=4 + n))
        style_axes(chart, "Política (s, Q)", "Ganancia anual promedio (COP)")
        chart.y_axis.numFmt = MILLIONS
        means = [p["mean"] for p in unif] + [e["mean"] for e in emp.values()]
        chart.y_axis.scaling.min = (min(means) // 1_000_000 - 1) * 1_000_000
        chart.x_axis.scaling.orientation = "maxMin"   # first policy on top
        legend_bottom(chart)
        chart.height, chart.width = 9.5, 17
        ws.add_chart(chart, f"J{max(last_row, 4 + n) + 3}")


def main():
    out = sys.argv[1]
    values = json.load(open(sys.argv[2], encoding="utf-8")) if len(sys.argv) > 2 else {}
    policy = values.get("recommended", [7, 15])

    wb = Workbook()
    build_parametros(wb, policy)
    reps = values.get("replications", {})
    build_simulation(wb, "Uniforme", "=RANDBETWEEN(3,7)", "uniforme", reps.get("Uniforme"))
    build_simulation(wb, "Empirica", "=VLOOKUP(RAND(),TablaEmpirica,2,TRUE)", "empírica",
                     reps.get("Empirica"))
    build_politicas(wb, values.get("policies", {}), policy)

    # "Automatic except data tables" (recommended by the statement); calcId of
    # current Excel so the stored replications are shown as saved.
    wb.calculation = CalcProperties(calcId=191029, calcMode="autoNoTable",
                                    fullCalcOnLoad=False)
    wb.properties.creator = "Juan Andrés Young Hoyos (1000337497)"
    wb.properties.title = "Taller 2 - Simulación Montecarlo de un sistema de inventarios"
    wb.properties.language = "es-CO"
    wb.save(out)


if __name__ == "__main__":
    main()
